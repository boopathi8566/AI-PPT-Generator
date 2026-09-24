from pptx import Presentation
from pptx.dml.color import RGBColor
THEMES = [
    (RGBColor(79, 70, 229), RGBColor(255, 255, 255)),   # Modern Indigo
    (RGBColor(2, 132, 199), RGBColor(255, 255, 255)),   # Ocean Blue
    (RGBColor(124, 58, 237), RGBColor(255, 255, 255)),  # Royal Purple
    (RGBColor(5, 150, 105), RGBColor(255, 255, 255)),   # Emerald
    (RGBColor(234, 88, 12), RGBColor(255, 255, 255)),   # Sunset Orange
    (RGBColor(13, 148, 136), RGBColor(255, 255, 255)),  # Teal
    (RGBColor(219, 39, 119), RGBColor(255, 255, 255)),  # Rose Pink
    (RGBColor(31, 41, 55), RGBColor(255, 255, 255)),    # Midnight
    (RGBColor(8, 145, 178), RGBColor(255, 255, 255)),   # Sky Cyan
    (RGBColor(30, 58, 138), RGBColor(255, 255, 255)),   # Royal Navy
    (RGBColor(244, 63, 94), RGBColor(255, 255, 255)),   # Coral
    (RGBColor(183, 121, 31), RGBColor(255, 255, 255)),  # Golden
    (RGBColor(22, 101, 52), RGBColor(255, 255, 255)),   # Forest
    (RGBColor(139, 92, 246), RGBColor(255, 255, 255)),  # Violet
    (RGBColor(185, 28, 28), RGBColor(255, 255, 255)),   # Crimson
    (RGBColor(71, 85, 105), RGBColor(255, 255, 255)),   # Slate
]
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def create_presentation(presentation_data, output_path, theme=0, design=0):
    prs = Presentation()
    



    # Remove default slide
    if len(prs.slides) > 0:
        slide = prs.slides[0]
        del prs.slides._sldIdLst[0]

    # Theme colors
    DARK = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    WHITE = RGBColor(255, 255, 255)
    LIGHT = RGBColor(241, 245, 249)

    for slide_data in presentation_data["slides"]:

        slide = prs.slides.add_slide(prs.slide_layouts[6])


                # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE
        # Top header bar
        header = slide.shapes.add_shape(
            1,
            0,
            0,
            prs.slide_width,
            Inches(0.25)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = THEMES[theme][0]
        header.line.fill.background()

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(0.6),
            Inches(12),
            Inches(1)
        )

        title = title_box.text_frame
        title.text = slide_data["title"]

        title.paragraphs[0].font.size = Pt(30)
        title.paragraphs[0].font.bold = True
        title.paragraphs[0].font.color.rgb = DARK

        # Content box
        content_box = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(1.8),
            Inches(11.5),
            Inches(4.8)
        )

        text_frame = content_box.text_frame
        text_frame.clear()

        for point in slide_data["content"]:

            paragraph = text_frame.add_paragraph()
            paragraph.text = point
            paragraph.level = 0

            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = DARK
            paragraph.space_after = Pt(12)

            paragraph.text = "• " + point

        # Footer
        footer = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(7.0),
            Inches(12),
            Inches(0.3)
        )

        footer.text_frame.text = "AI PPT Generator • Powered by Gemini AI"

        footer.text_frame.paragraphs[0].font.size = Pt(10)
        footer.text_frame.paragraphs[0].font.color.rgb = BLUE

    prs.save(output_path)

    return output_path